import base64
from io import BytesIO

import openpyxl
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestExportRegression(TransactionCase):

    def setUp(self):
        super().setUp()
        self.company = self.env.company
        self.partner = self.env['res.partner'].create({
            'name': 'Proveedor Test',
            'vat': 'J-31527189-4',
        })
        self.company_id = self.company.id
        self.journal_purchase = self.env['account.journal'].search([
            ('type', '=', 'purchase'),
            ('company_id', '=', self.company_id),
        ], limit=1)
        self.account_expense = self.env['account.account'].search([
            ('account_type', '=', 'expense'),
            ('company_ids', 'in', [self.company_id]),
        ], limit=1)
        self.tax_general = self.env['account.tax'].search([
            ('amount_type', '=', 'percent'),
            ('amount', '=', 16.0),
            ('l10n_ve_tax_type', '=', 'general'),
            ('type_tax_use', '=', 'purchase'),
            ('company_id', '=', self.company_id),
        ], limit=1)
        if not self.tax_general:
            self.tax_general = self.env['account.tax'].create({
                'name': 'IVA General 16%',
                'amount_type': 'percent',
                'amount': 16.0,
                'l10n_ve_tax_type': 'general',
                'type_tax_use': 'purchase',
                'company_id': self.company_id,
            })

    def _create_vat_book_line(self, period_month=None, **kwargs):
        """Helper para crear vat.book.line."""
        if period_month is None:
            period_month = '2026-09'
        defaults = {
            'book_type': 'purchase',
            'period_month': period_month,
            'partner_id': self.partner.id,
            'invoice_number': 'FAC-001',
            'control_number': 'CTRL-001',
            'operation_code': '33',
            'base_general': 1000.0,
            'vat_general': 160.0,
            'company_id': self.company_id,
        }
        defaults.update(kwargs)
        return self.env['l10n.ve.vat.book.line'].create(defaults)

    def test_export_xlsx_includes_odoo_lines(self):
        """Verifica que el export XLSX incluye líneas source='odoo_invoice' y totales cuadran."""
        # Usar un período único para evitar conflictos con datos existentes
        test_period = '2026-09'

        # 1. Crear línea source='excel_import' (línea base)
        self._create_vat_book_line(
            source='excel_import',
            invoice_number='FAC-EXCEL-001',
            base_general=500.0,
            vat_general=80.0,
            period_month=test_period,
        )

        # 2. Crear línea source='odoo_invoice' (la que se prueba)
        self._create_vat_book_line(
            source='odoo_invoice',
            invoice_number='BILL/2026/09/0002',
            base_general=1000.0,
            vat_general=160.0,
            period_month=test_period,
        )

        # 3. Crear línea source='manual' (otro origen)
        self._create_vat_book_line(
            source='manual',
            invoice_number='FAC-MAN-001',
            base_general=200.0,
            vat_general=32.0,
            period_month=test_period,
        )

        # Ejecutar export
        wizard = self.env['l10n.ve.vat.book.export.wizard'].create({
            'period_month': test_period,
            'company_id': self.company_id,
            'book_type': 'purchase',
        })
        wizard.action_export_xlsx()

        # Decodificar el archivo
        xlsx_content = base64.b64decode(wizard.template_file)
        wb = openpyxl.load_workbook(BytesIO(xlsx_content))
        ws = wb.active

        # Verificar que la hoja tiene datos
        # Contar filas de datos (excluyendo header y total)
        data_rows = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is not None and row[0] != 'TOTAL GENERAL':
                data_rows += 1

        # Debe haber 3 líneas de datos (excel_import + odoo_invoice + manual)
        self.assertEqual(data_rows, 3, "Debe haber 3 líneas de datos en el export")

        # Verificar que la línea odoo_invoice está presente
        found_odoo = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[3] == 'BILL/2026/09/0002':  # invoice_number
                found_odoo = True
                break

        self.assertTrue(found_odoo, "Línea odoo_invoice no encontrada en export")

        # Verificar totales
        # base_general total = 500 + 1000 + 200 = 1700 (columna 6)
        # vat_general total = 80 + 160 + 32 = 272 (columna 7)
        total_row = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == 'TOTAL GENERAL':
                total_row = row
                break

        self.assertIsNotNone(total_row, "Debe existir fila TOTAL GENERAL")
        self.assertEqual(total_row[6], 1700.0, "Total base_general debe ser 1700")
        self.assertEqual(total_row[7], 272.0, "Total vat_general debe ser 272")
